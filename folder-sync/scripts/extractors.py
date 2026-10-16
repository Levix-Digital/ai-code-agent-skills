import os
import re
from pathlib import Path

SUPPORTED_EXTENSIONS = {
    '.docx', '.pdf', '.md', '.txt', '.csv', '.json', '.yaml', '.yml', 
    '.xml', '.html', '.htm', '.log', '.py', '.js', '.ts', '.tsx', 
    '.jsx', '.java', '.go', '.rs', '.sql', '.sh', '.bat', '.ps1',
    '.c', '.cpp', '.h', '.hpp', '.cs', '.rb', '.php'
}

CLOUD_SHORTCUT_EXTENSIONS = {'.gdoc', '.gsheet', '.gslides'}

IGNORE_DIRS = {
    '.git', '.svn', '.hg', 'node_modules', '.venv', 'venv', '__pycache__',
    '.folder_brain', '.chroma_db', '.next', 'dist', 'build', '.idea', '.vscode'
}

def is_ignored(path: Path, root_path: Path) -> bool:
    try:
        rel = path.relative_to(root_path)
        for part in rel.parts:
            if part in IGNORE_DIRS or part.startswith('.'):
                if part != '.':
                    return True
    except Exception:
        pass
    return False

def extract_text_from_file(file_path: Path) -> dict:
    ext = file_path.suffix.lower()
    
    if ext in CLOUD_SHORTCUT_EXTENSIONS:
        return {
            "type": "cloud_shortcut",
            "content": f"[Google Cloud Document Pointer: {file_path.name}] - Necessário exportar como .docx ou .pdf para indexação.",
            "sections": [],
            "word_count": 0
        }
    
    if ext not in SUPPORTED_EXTENSIONS:
        return None
    
    content = ""
    sections = []
    
    try:
        if ext == '.docx':
            import docx
            doc = docx.Document(str(file_path))
            full_text = []
            curr_heading = "Geral"
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                if para.style.name.startswith('Heading') or para.style.name in ['Title', 'Subtitle']:
                    curr_heading = text
                full_text.append(text)
                sections.append({"heading": curr_heading, "text": text})
            
            for table in doc.tables:
                table_rows = []
                for row in table.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        table_rows.append(row_text)
                if table_rows:
                    t_text = "\n[Tabela]:\n" + "\n".join(table_rows)
                    full_text.append(t_text)
                    sections.append({"heading": "Tabela", "text": t_text})
            
            content = "\n\n".join(full_text)
            
        elif ext == '.pdf':
            import pypdf
            reader = pypdf.PdfReader(str(file_path))
            full_text = []
            for idx, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    full_text.append(f"[Página {idx+1}]\n{page_text.strip()}")
                    sections.append({"heading": f"Página {idx+1}", "text": page_text.strip(), "page": idx+1})
            content = "\n\n".join(full_text)
            
        else: # Text, Markdown, Code, JSON, CSV
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
            
            # Simple markdown heading parsing
            lines = content.splitlines()
            curr_h = "Geral"
            buffer = []
            for line in lines:
                if line.startswith('#'):
                    if buffer:
                        sections.append({"heading": curr_h, "text": "\n".join(buffer)})
                        buffer = []
                    curr_h = line.lstrip('#').strip()
                else:
                    buffer.append(line)
            if buffer:
                sections.append({"heading": curr_h, "text": "\n".join(buffer)})
                
    except Exception as e:
        return {
            "type": "error",
            "error": str(e),
            "content": "",
            "sections": [],
            "word_count": 0
        }
        
    words = len(content.split())
    return {
        "type": "document",
        "content": content,
        "sections": sections if sections else [{"heading": "Geral", "text": content}],
        "word_count": words
    }

def chunk_document(doc_info: dict, target_chunk_size: int = 800, overlap: int = 150) -> list:
    content = doc_info.get("content", "")
    if not content:
        return []
    
    sections = doc_info.get("sections", [])
    chunks = []
    
    # If sections exist, chunk respect section boundaries
    for sec in sections:
        heading = sec.get("heading", "Geral")
        text = sec.get("text", "")
        page = sec.get("page", None)
        
        if len(text) <= target_chunk_size:
            chunks.append({
                "heading": heading,
                "text": text,
                "page": page
            })
        else:
            # Sub-split long section with overlap
            start = 0
            while start < len(text):
                end = start + target_chunk_size
                # Try finding a paragraph or sentence break
                if end < len(text):
                    break_point = text.rfind('\n', start, end)
                    if break_point == -1 or break_point <= start + target_chunk_size // 2:
                        break_point = text.rfind('. ', start, end)
                    if break_point != -1 and break_point > start + target_chunk_size // 2:
                        end = break_point + 1
                
                chunk_text = text[start:end].strip()
                if chunk_text:
                    chunks.append({
                        "heading": heading,
                        "text": chunk_text,
                        "page": page
                    })
                start = end - overlap if end < len(text) else len(text)
                
    return chunks
